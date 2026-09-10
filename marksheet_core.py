"""
Core logic for the Mark Sheet Generator:
- parse_marks_docx: reads the uploaded .docx (Sl.No, Name, APAAR ID, Marks table) into a DataFrame
- compute_marks: scales raw marks to /25 and derives SA (SA component) marks from the /25 score
- generate_marksheet: builds the final formatted mark sheet .docx matching the college template

Note: the college logo is embedded below as a base64 string (LOGO_BASE64) so this
single file has no dependency on an external assets/ folder or image file.

IMPORTANT: this file must be uploaded to GitHub as a direct file download
(GitHub "Add file -> Upload files"), not copy-pasted by hand. The embedded
LOGO_BASE64 string is very long and manual copy-paste easily truncates it,
which silently disables the logo (no crash, it just won't show).
"""

import io
import base64
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------------------
# Embedded college logo (base64-encoded JPEG). Replace this string with your
# own logo's base64 data if you want to change the crest shown on the mark sheet.
# ---------------------------------------------------------------------------
LOGO_BASE64 = (
    "/9j/4AAQSkZJRgABAQEAYABgAAD/2wBDAAMCAgMCAgMDAwMEAwMEBQgFBQQEBQoHBwYIDAoMDAsKCwsNDhIQDQ4RDgsLEBYQERMU"
    "FRUVDA8XGBYUGBIUFRT/2wBDAQMEBAUEBQkFBQkUDQsNFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQUFBQU"
    "FBQUFBQUFBT/wAARCADHAO4DASIAAhEBAxEB/8QAHwAAAQUBAQEBAQEAAAAAAAAAAAECAwQFBgcICQoL/8QAtRAAAgEDAwIEAwUF"
    "BAQAAAF9AQIDAAQRBRIhMUEGE1FhByJxFDKBkaEII0KxwRVS0fAkM2JyggkKFhcYGRolJicoKSo0NTY3ODk6Q0RFRkdISUpTVFVW"
    "V1hZWmNkZWZnaGlqc3R1dnd4eXqDhIWGh4iJipKTlJWWl5iZmqKjpKWmp6ipqrKztLW2t7i5usLDxMXGx8jJytLT1NXW19jZ2uHi"
    "4+Tl5ufo6erx8vP09fb3+Pn6/8QAHwEAAwEBAQEBAQEBAQAAAAAAAAECAwQFBgcICQoL/8QAtREAAgECBAQDBAcFBAQAAQJ3AAEC"
    "AxEEBSExBhJBUQdhcRMiMoEIFEKRobHBCSMzUvAVYnLRChYkNOEl8RcYGRomJygpKjU2Nzg5OkNERUZHSElKU1RVVldYWVpjZGVm"
    "Z2hpanN0dXZ3eHl6goOEhYaHiImKkpOUlZaXmJmaoqOkpaanqKmqsrO0tba3uLm6wsPExcbHyMnK0tPU1dbX2Nna4uPk5ebn6Onq"
    "8vP09fb3+Pn6/9oADAMBAAIRAxEAPwD9U6KKKACiiigAooooAKKKq317BYWstzczR29vGu6SaVtqqvruoAtUV4J4y/bk+BXgO5Fp"
    "qHxJ0e8v92xbPRWfVJmb+7stlfDexrl0/bU1vxb8nw++AfxM8Ts3MV3q2nxaLYy/7s87/wDstAH1HRXy4njv9rHxTCJdP+Gfw78A"
    "R/xJ4q8RT6i6r6/6Gm3/AMerN1Lwv+0deSSp4h/aC8EeBJIo0mkg0PwxFceWrPsVv9Ml+6zjarH+Lj2oA+taK+TNJ+A3jfxZrusa"
    "JqP7VvizU9Y0gwrqVlodnp2ny2vmrvi3qkTbN6/MtcNb/DXwDqn9pyT/ALUfxvmexLefv1+4skZVuEtXeLbap5qJOyxM8W5VZvmo"
    "A+7KK+Htd+Bnwx8OwJLq/wC0J8YoR/b6eFvMl8XXn/IRYbhD/qv7rbt/3F/vVBefBv4a6P4qv/DMnx++Oun6nY3UlrcS/wDCR6il"
    "p56WX2xovtX2fyGkW3zLs37qAPuiivhfwv8ADvwT4iVpvDP7V/xelVdDXxQxvdf85YtO3MPPfz7X5F+X7rfNx92uth+FPxI0vwbH"
    "4r0X9rvULTwrNbJexat4g0HS7y0+zsuVdpX2fL8y/wAS0AfXdFfJel2/7Tttp9nf+HPi18KPiLp16M2V5rOkz2SXX+41nK6t/wAB"
    "3Vpr8XP2nPCrbNe+BGh+Loh9+88I+LYoAP8AdhukVm/76oA+oaK+XP8Ahvbw94bk2fEP4c/Eb4aIv+svta8OS3Fjn/ZntvN3f981"
    "6b8Ov2ovhH8WVhj8KfEXw/rF3N9yxS+SK7P/AGwfbL/47QB6tRRRQAUUUUAFFFFABRRRQAUUUUAFFFFABXyb+2x8RviPpfir4V/D"
    "f4eeILfwVeeO7y9t38UTw+f5EkESSxQKv8JlZtu6vrKvlH/go54evV+ANt480dT/AG78O9esPFVps+8fKl2Sr/u7JWb/AIBQB8yf"
    "Fz/goh42tvCvgfwvL4kPwx+IGm6q2jfESKPQvtt1ZorIn2+13r5Bif522ff+dNvy/M3p/wAVPgH8MPh9/wAIxqnxPuPHnx0TVma4"
    "XVvEWvSvplmi+V0WJkiTcsm5Eb5W2P8AN8tcJ+0p8QPAnxO+M3h/xN8GdP1L4peLdRsl0bxv4b8PWMtxZanok8X3bq4QeXFOnyhW"
    "/hZU37fKWtPwj8SrL4E+HdV+A3xe8H+NvHNz4Pm+0eE7jwzZzy3V/o8kb+Ud8Drt8pN8T/Ps2/J82ygD0KX4zeHfhZr3xJ8GfC7w"
    "b4V8MS2OnL/wi2t6Ppe621C+awgv0t52iVU/frP+6+f59j/7G67qHjn49/FLQW03R9G1Lw34lXxNa3dnqX+o0w6c2lPOn+keVmW2"
    "a8iXcjL5uyZYn2vVb4RfHT4sfGzwRHqvwL+EHgrwP4OupdkGs+KNWVkl8jbb/NZWS70dViRBvb7sSfw7K7k/s9/Hrxyrf8Jr+0Te"
    "aPay/wCs0zwJoNvYeX/uXUm+WgDG1z9n7xd410vx1Kug2fhnUPGmi6TaXf8AbF8dQSxl+2y/2vbrsl3NbtbrEyKrru+7+6/g84+I"
    "Pwl8OLptpD4z/aB8H+HfE2j6E3h7Ttf/ALQiinnt0uElgS4gll+Zdm5WTezbkil37q9kX/gnn8K9WVH8ZXfjD4lXCHeJvF/ii8uD"
    "/wB8o6L/AOO13fhz9j74H+Eo1/sz4T+EUdfuyXOkxXEo/wCByqzfrQB82+H/AI3fs+/C34xap4/m/aCtdS1fVru+k1Kysbc3FrPb"
    "yP8AuIl+zxM+6FEt1V3d/uPtVfN+Xk3+LH7OXiWw1Ozu/iJ4o8WaPMuqW1jY2PhG88qzivtQiv7j5vsv7397FEvz/wAK/d+av0H0"
    "nwjoXh1VXStE07TkX7q2dokQ/wDHVrboA/N/W/En7NuvQ3unDRfihbaDLcX91Houh+HL/TbS1lu4reKV0SCKL7qQfKrbl/ey7lbd"
    "Vi++IXwCbxFqXiOd/iVaeKb5bqO/8QXHhGf7RPFPYRWbK3+i/c2RLKv8Svu/gdkr9GKKAPzE0vxd+yr4Z/t6x/4W14m0K212Zv7Q"
    "ttQ8P38TeR56Tvao7Wu2KL91Emz+5uX+OvUtF+LnwPvPhJoXhDQvj34Xmfw/rq65pkviV/IhZUu3nhtbhXZNyoHCqybdnlo235Np"
    "+58blw1cxr3wy8H+KE2614T0PVg3a+02Gf8A9CWgD5W8QfC22+NENzD4a1/4XeM59X8PNo7XljdLEmgXUtxPcT6hp8ESz75X+0I3"
    "+tibfbxNv+b5ew1C3+Jvh3wX4zsIfDviL+1tQ8XrfnWrC7tZS2lvqESv5CrL5vmrYQ7duxfmaun8UfsJ/AHxgG+3fCfw3Bu/i0y1"
    "+wN/5A2VzB/YL8NaCwbwH8RPiZ8OwvMVtonimWW09t0Vx5u5fagBdc+LfjLVP2hUsbG51TQvBVvplnOun3Vj9nuNRm2S3NxsWeyb"
    "zP3XlRbEuInR0f5X/h8W8UeOPhl8UvDupa18QPhr4F8Z3F1pml38FnollPFqFvJqMuyG3a82fvfl812liddv2eX91/FXszfDH9qH"
    "wHk+GfjJ4X+IMS/csfHXh42jIvp9os23Of8AaZa8i+If7QD/AA48S6L4N+NP7PLW+p61ePqWnaj8MdTW7uL26hVle6iii8q5iZEl"
    "b592/wCZ6AKnjr4d6d+z7+zzqnxc+HvxQ8cfCmztbZm07w3LrEHiDSrmXcUgjiilaVW819vzbtyp95V2tWH4d/a++J/7Rvw58I+B"
    "/A3i2x0jx7Fp8ut+PvGkOlSxWnh22iZmW38qVPnnb5FfZ8vytt+XcyQ/E34ieJv2nNdk8U+BPBXiPxH8PvhHawTaRoep20puda8R"
    "OqrG86SnzZVtVbcwyzMyfxebXI2vxc8F/Dr9if4heGtF1nULz4/+MpY7LxPYeILSWz1i71G/fy5R5Uo3NGkUsqrt/wB75WegD7t/"
    "Y3+LmvfHT9m7wV438TWcVrrmqQS/avJXakpinli81V/h3+Vu2/7Ve31x/wAJ/Att8L/hn4U8H2RV7fQdMt9PWRf4/KjVN3/Asbvx"
    "rsKACiiigAooooAKKKKACiiigArnfHvg+y+IXgfX/C+ojOna3p8+nXW3r5UsTI36NXRUUAfDfwK0n43/ALHngO2+HVn8E9N+Img2"
    "M8zQeJfCuu29hNeb2Z91xbz/ADeZ/Du3fwotd1e+OP2pPidDJpmh/DXQfg3DL8reIPEuuRaxcIndorW2XaJP+urba+qqKAPKv2d/"
    "gfo/7Ovww03wZpN1daoYJZbm91O8OJr26lfdLK393PGF/uqv3vvV6rRRQAUV518WPj94A+BumrfeOPFWn6CknMFvK++6uR0/dQJu"
    "kl5/urXk8f7QPxg+KjlfhV8HZ9J0xuI/E3xMnbS7fPUMtlHvuJVbs3yUAfTtFfLepfBP4ta9Zy33xL/aMvPD2mYDyWPgjTrXRobf"
    "/t6nEsrD/e21xH/Cj/2Zde1D7Bq/jPxL8YNZSL7Q0LeKtU1+4Me/7zRWbt8u7/ZoA+zbjVbKzOJ7u3hP/TWVVqa1vbe8XdBNHMP+"
    "mbhq+RLP9mf9l21j2W/wT1KYf9NvCGsn/wBGxVy198I/2O7rW9Y0d/htqWnaxo9st3fwWvhrXraS0gfdslZoohtT5H+f7vytQB92"
    "UV8SaL8JfgJeWNjceBfjf4u+HMV6iXFjHp3jm6t0uFb7rLDfO/mr/wABrvk+Hn7RngmNbjwj8XfD/wAStNx+603xxoy2zlP9m9s+"
    "Xb/eioA+naK+Xv8AhsbUfhuwh+N/wy8RfDaEMqt4isf+JzoWP7zXUC7otx/hdK+hPB/jTQfH2h2+teG9Zsde0q4/1V7p1ws8T/Rl"
    "4oA3a+ff2hP2c9Y+I3i7wp8RvAXiSHwr8TvCqyx6fdX1ubixvLeX79rdJ97YefnT5l3N/slfoKigD5dj+OH7R3hz/RfEH7OcPiCa"
    "L5TqnhXxbai3nb1WKfbKi/71cLqHwR+JH7Unx8+HPxA+I/w40b4Y+H/A119vt7U6nFqWr6nKArxI8sX7tYklRG2N/t/3q+3KKACi"
    "iigAooooAKKKKACiiigAooooAKKKKACiiuS+I3xE0H4U+DNW8V+KNRg0nRdLgNxcXEzdB/dXn5mY4VV6s2KANfXte03wzo15rGsa"
    "hb6VpdnE01ze3cqxQwqvVnZvu18w/wDC6PiX+1NJNZfBKAeCfAG9opvibr9nulvB/F/Zdm3+t/66y7V+9/EtZfg/4c+KP20NVsvH"
    "PxXsbjQfhXDKt14b+G9wBvvlz8l7qn97d95IPu9P+B/X1raw2NvHBBEkMMa7Y44kCqq4+6tAHjvwj/ZQ8A/CPUJNdjtbnxR43nO6"
    "58YeJpjf6rM+Of3r/wCr/wB2LaOK9fvrVb+0nt2kkhWWNoy8LbXXdxlW7NV2vNvi58fvAfwN0+O58Y+JLbS5bg/6Jp67pby7bpti"
    "gQNI/P8AdWgD50/aM+FOnr4p8M6efBmtWMEV3HfwfFCzbUvEOp2LxPEz2/lRLLP+9TdErSt5W3f8v3VrvfD3w68St+0t4Z+Ktne6"
    "xrmhax4c1DQ9TbWLGDTZdMiW6insgtu0UU53N9o++rsN38O6q0Xxm+O/xah3fDj4T23gbR5OYtd+KF09vKy5/h0633Sr/s73WrUf"
    "7O/xi8V5fxn+0VrtvG3LWPgrRrPSkiPosrpLK3/AqAKfgX9knTb3wHq2gfE/QdD1zxLMksK/EC3lafWLxpGfbdebLFvtZ0+QqqO6"
    "o33ePlrO8Xfs3/ELw/8AC3x5YeCtT8I6h4j1vQpdKa/uNBuF1K+XymRInupb9lX77/wbFZvuV0K/sT6dNGv274vfGLUnX+K58b3C"
    "f+OxBFpD+xeLcb9L+OHxk0tl+6v/AAlrXSL/AMBniegCHx/4S+LHgPw34T8O/Dy1fUPDOjeGP7NTT7CKwe4lv4kRLf7R9sdU+zbE"
    "+byn37v71eDeH/gv4z+C8fwY+EGj3thpXi7XpW1TWvE3hma4tb61s4GSe/adNzwT7nl8hJZfl+ZNsX9z3uT4R/tF+B1M3hP426Z4"
    "zRB+70jx94eiVG/3rqz2Sf8AjjV5h4w+Iur+C/iBZeOvi38OfEvwz8UWVmunRfEPwjdy69oD2au8rxXkCfNFAzv/ABxb/wC66bFa"
    "gD7L0TRptN0WGwv9TuNdlVWD3V9FEskyn+8sSIn3fl4WvBfGX7HOlWGuXHi74O61c/B/xs53Svose7StQYEfJeWH+qdf9pQrclvm"
    "r3Hwb488PfELQbbXfDGtWGv6RPxFfafcLNE2P4dy/wAX+zXSUAfNHgX9qLVvDPiyw8AfHXRLfwD4yum8rTNes5Gfw/rrf9Os7/6q"
    "T/plL83T+8Fr6Xrj/iN8NfDPxa8H3vhjxdo1vr2iXiFZrS6TK57Ov8SOv8LKdymvmzQ/G/if9ivxTpng74i6pd+J/gxqlwtn4e8c"
    "Xx3XGiSMfkstRb+KL+FJ+38Xy/cAPsOiokkWRFZW3K33WWpaACiiigAooooAKKKKACiiigAooooAKKKKAInkWNGZm2qv3mavjXwr"
    "Zv8AtzfGBvFmrL53wH8E6g0Xh3T5F/deJdUi+WW9lX+O3ib5UX7rt/wNa7D9tzxdq+q6P4V+DHhK8az8WfE68fTWuo/vWOmRrvv5"
    "/wDv18v/AANq968C+BtG+HHg/R/C/h+1Fjo2lWyWlrbqSdqKMf8AfX+1QB0tFFfCP7cf7WWg2XjrTfgXbeO4vAcWpL5ni3xQoZ5d"
    "OsmXP2WDYjfv5VP/AAFWX+98oB6J40/aG8YfGrxnqXw7/Z/+ytNp0pg8QfEbUIvO03Rm/wCeVun3bq6/2PuLld38Wzvfgv8Asr+D"
    "fg1qEuurHd+LfHV58194x8SSfatSuG4DbXb/AFSf7Kbf+BV5j8N/2xv2VvhH4L0zwp4V8cabpOiadEIYLeHTrznvvb918zNyzN/F"
    "XV/8PEv2dv8Aoptl/wCAN5/8ZoA+kKK+b/8Ah4l+zt/0U2y/8Abz/wCM0f8ADxL9nb/optl/4A3n/wAZoA+kKK+b/wDh4l+zt/0U"
    "2y/8Abz/AOM0f8PEv2dv+im2X/gDef8AxmgD6QpuNy4avnH/AIeJfs7f9FNsv/AG8/8AjNH/AA8S/Z2/6KbZf+AN5/8AGaAK/j79"
    "kf8AsnxHceOfgprCfCzx+3z3MNtEP7E1j/YvLNfl7/61BvXcW+Zq3fgT+0sPH/iC9+H/AI30Z/A3xZ0iMPe+H7ht8d5F/wA/dlL/"
    "AMt4G/76T+L+8cr/AIeJfs7f9FNsv/AG8/8AjNeO/tHftFfs1fHLw3bT2fxdtfC/jzQ5TeeHfFVnYXXn6fcqf+uPzRPgKyfxLQB9"
    "5VgeMvBujfEDwtqnhzxBYRanoupQtbXdpMp2yI3b/wCv7V4h+xR+1Xp/7Unwya8lmtl8X6JILLXbW1Y+UZf4biL/AKZS7Sy/8DX+"
    "Gvo6gD5K/Z78U65+z98UH/Z68b6lLqNg8LXvw/8AEN4ctf6en37GVv8AnvAP++l/u/KD9a14N+118Fr34yfCiWfw25s/H/hi4XXv"
    "DF/CfnjvoPmVP92VRs/4Ep/hrqf2dPjLZfH34N+GfHNkiQNqVt/pdrzm2uU+SeLk/wAMisP93bQB6hRRRQAUUUUAFFFFABRRRQAU"
    "UUUAFFFVri6js7aWeVtkcSM7N6KvWgD5V+CMP/C3/wBs74yfEacGXTfBiweAtEf+ESJ+/wBQ/wCBCVlX/davrOvlT/gmnbSXP7LO"
    "neJ7pcaj4v1nVNfvG/vSy3cqf+gxJX1XQB578ePitYfAz4P+LfHmoqskGiWL3CxM23z5fuxRf8DlZF/4FXy74T/YqvvEP7MNrqOp"
    "XraZ8eNS1AeO/wDhKGTZcWutv+8SJv7sSrtiZPu/eau0/bojXxjqnwO+GMm57XxZ45tZtRhAwJ7GzVriaI/+Q/8AvmvrCgDxf9mL"
    "49D49fDv7ZqFk2heNNFun0jxLoMv+ssb+L5ZVK9kb7y//YtXtFfIX7SGg6n+zl8UrX9ovwnazXWjNHHpvxB0e1Xe15pufkv1XvLb"
    "/wDoH91d9QfF74seNYdfuPEOj+Ir3TfCV3bwXXhTXLdrVPDzRvaq6PfyvE8srTXR8jyk2vseJokZ97IAfX0is0beWwV9vyuw3V8n"
    "6l8ffi5peialIi+G9Z1SPx0ng+2t7DQJo3ZPK3vcbJtRRWZv4V81FH95q80/ab1q7n+IHhn4paR8XPHOm/Be/u10LxPD4X1ZrWTw"
    "9ebdsVw0UsTeVEdy+ajqrLuVv49tek+PP2e9J+GvgXVfF+v/ALRnxqtfD+l232y6u08RxSjy/wDZVLX5v+A0AZTftmeI9K+H9xqu"
    "rW3h+w1qXwlrGpadaXSS27XWs2t08EVh5LTbvNYGLfaqzOrttV3+9XY/Ej4zfEvwFqXxXlWTw3d6d4R8HReJ7S0Oi3K3M8sxv1SB"
    "5ftm3EX2JdzIvzb/AOCvzh+Nn7amgaTb3Nn8LPi18c9e1JW2x6lr3iCC3svvfeWNYPNf5f72yu9/ZK/aA0L436tpfhDxx8bfjN4S"
    "8bXrPFFcQ+KIBpkxVHbO+SDdG2F27W3ZY9aAP0m+E/xS1vx9408eaTq/hufw1DoE9nHa2l60T3bLLb+azStFLLE3zH5djdua9Xr4"
    "9+K3wH0v4N/DfXPG/iL9on40W2i6TavPK6+KIN8v9xE/0X5mZtqL/vVc/Y1tfFnw98C6NdfF7xd4kv8AxR8QLxm0nSvENzLdf2dE"
    "kTyw2rPt2pO0W933bN23bt+SgD62r5W/aw8e658QPE2k/s9fD2/Nn4p8VQG48Q6xB10LQ87ZZf8ArrL/AKpP97+HejV6x+0B8cNL"
    "/Z/+GOo+KtRhkvbvzFstK0uDJl1G+l+WC3jX1ZvToqs3auR/ZP8Agfq/w18Pax4r8cTpqPxS8aXA1TxHfr0jbH7qyiHaKBDsXqPv"
    "Y+XbgA8i+Mvw60D9i7xp8KPip4Lsl0XwdpfkeCfFlrD92TTJ2/0e7k4+9FOdzP8AebetfbleW/tNeAYfid+z38RfDFwgmOoaHdLA"
    "D/DOsZeF/wDgMqI34Vm/shePpPiZ+zD8MvEdxMZrq60O3iuZmP354l8qVv8AvuN6APZK+Sv2b0Hwk/aq+NvwnRDDo+rNb+PdDi/h"
    "WO5/dXu3+6ouFUKK+ta+UvjVIPCP7ev7POtxfIPEmla94evH9UiiS6iX/v7QB9W0UUUAFFFFABRRRQAUUUUAFFFFABXP+PldvA/i"
    "JY/9Y2nXAX6+U1dBUMkazRtG67lZdrLQB86f8E7Gik/Yt+FZjxt/s6Rf+BfaJd365r6Rr4//AGAtbuvBv7LOt+GP7Ou9a1X4d69r"
    "ehSababftFxLBcNOsSb2Vd7ecqrvZRW/8QPix+0NqXgvWrjwp8GIvC17bQG7trnWvEFtdTv5Z3+UtnapL5juq7dnmp9/7y0AU/2k"
    "WCftofsmyyc24u/EyMzfd3NpqbK+ra/Pj9qzWvFnwz8ReBfEHi3WofFf/CvPEWj+Jpdbh05LCZdMvJZrO9ieJX2/K6W+1/8Aprtf"
    "7m5/0CSRZEVlbcrfdZaAK1/p9tqVnPaXkEdzbTxtHJBMu9ZFbhlZT94c18OeGfBC/Av4sWXwG8Q6vqGmeCtS1M+JPhZ4iidWbTLx"
    "d3m6W3mqyPt819ivu3K/9512/eNeVftFfA3Tvj98M73w7dTnTtUikS/0bV48iTTb+I7oLhSP7rf+OswoA+bvAni2OTVvEPhLxZe6"
    "L4tt/GmuX9p4r8P6o7trtpaxRPb/AG+9eLbBbQeVaxbU8qJdjptld/veZ+IvhP4/8ZabJ+yNffEhfDdlbKNT8M6rfWaXKeJPD4+5"
    "as27Pn2rKvyL95U3fdVXr0LQNX134/eF5LnxNa2tx8Tvhxu0PxR4F1Wwn1XT7q63q1vqUVlFLEsjNhmRn+X52+ZNm+srwzpMX7Rn"
    "wlh8JeHLk+Hfif8AC+5eXwZqr6hDcXjSW2xJXuEg3xW9tK58jy/NdWWL5dyoNwB84/tOf8EqdM/Z/wD2f9f8e2/xBuda1LRFikmt"
    "ZNMWCKdWmSL5f3rMn3938XSvQtV/4Ihwt4ftG074qSJraopuVutI3W7t/Fs2y7l/8er1v9oz462/x7/4Ju/EzWJrJtH8T6clvpfi"
    "DQ5D+907UIr6BZYm/wBnPzL/ALLV6r+1l8Stf17WtE+BXw6vPI8eeM4nfUNShY/8SLR1Oy4vHx9xm+aKP/a7httAHhvwZ8F3Xxdv"
    "NC0XXvEN943+C3wMHlLqMOmM/wDwlOsW6fLst4t7TQ2qYRFXe0rf399eq+A/Dun/ABovrjxJ8QNA8M+O9MvriS6j8XaFqqr/AMI1"
    "5Sb1spVfypYHt/u+an73e250i/hxviBdp8K/BXw8+H3w88I6drfwzZ7fS9H13T/Edxpstvq32poN0t1AjeVuZ5X835lldZU+86K/"
    "N+MvCupfEL4gT/ATQtZ1C517xFHbav8AFnxM93FPLZ6ckSJDpqSxW9uu+VPl/wBUrbH3tv3vQB2PwP0WX9qn4yW3xf1SO4f4Y+C3"
    "k0z4e2d47y/2jKuY7jV5d/zO3y7Imbn5d3Drub7KrH8OeGtM8H+H9N0PRbKHTtI02BLW0tLddqQxKu1UWtigCnqUkcOn3UsrKsKx"
    "szM393bzXyF+wN420z4a/sA/DTWfENz9j08TvbNcMNqRm51iWCJnZuAu6VNzfwr9K9m/bA+IkXwq/Zl+I/iOSTy5odGnt7Uj/n5n"
    "XyIP/IsiV4vqXw/1fwH+zj8DPhRYeEbHxjrNvFa32r+G9Su0tbaeK0i8268x3jZdq3Utvwy/NuFAH1Ovjrw6+k3Wpprmmzadawtc"
    "T3kN3G8UUaruZ2bd933r5c+PmvweLPjh+yVq1nBc2seo6/f3ltHcKFlMBsy+5lPK7k2Ns+983zfNxXF/st/Bz4R/G/Q7yy8ZfADT"
    "dB168jbxZBcSwWrRT6dqN3dS2XlSwPuVEiQR7GVPufdr0b4iafb6/wDt8fAjwtYxhbbwN4W1fX5Yuu2GdVsIf/HloA+taKKKACii"
    "igAooooAKKKKACiiigAooooA+Sfhyx+C/wC3h8QfCUyeToPxO02LxXpHHyfb4F8q9iX/AG3Uea3/AAGuU8ffDf4hePfFnjXWPEHh"
    "Sbxx4UvPEscb+DtNMVrPFBpxdLR3M8sSXEV1FK8rtv8Akb7PtVtrLXqf7anwx1zxZ4B0nxr4It/M+I3w9v18QaGqKd9yq/8AH1af"
    "7ssX8P8AEyItX7XSvCn7anw28FeMbbxP4ks/Cd5A81zoWias1hHes2Fkt7xosSt5TI67VdOd1AHn/wAIPhXafEvwz8QfBnjvxAmq"
    "+KZPC2neE9Y0dZPtEukwbbmWL/SC7efL/pBVpfl/e2jf71dN+xF8SdU1XwTqHwx8ZS+X8Q/hvMuhapE7/NcwKP8ARLxO7JLEF+Y/"
    "3f8Aary/4VeJl0LTfDnjDwF8E4/AGhLPf29j/ZM8Uv8AbNnEzfaLe/VUVoLr/RWliaVmTfF5Tyr5vzdx8ZPBd/4vk8L/ALRPwJmt"
    "tW8Y2Fhl7GKXFv4p0hjuazc/89VPKE/db5ey7QD6xorzf4G/HPwx8f8AwTF4i8NTsAj/AGe+026Xbd6ddL9+3nj/AIHU16RQB8pf"
    "tVeA9a+GvizSv2h/AlnLfa94btza+KdFtVwdc0TO6VP+usP+sQ/7Hfaq15ZrXhvwzoum2vjvSPFd7p/gCSzuvEPgvTvCdssUSRRW"
    "D3TIlrEEiW6gktX/AHtxuWVbhon+avv7G5cNXwbrXw50X9n34lXPwt17S9Nuvgv8RruWfwhNq1ol1beG/EDq2612P92KXduRf95f"
    "79AHln7bngzXfF3wt8YfFz4Y6XdW1p4oVPDvj3wz5Yl8+S2uk+z6hH5TMkjI6Km9d3yS/wC/Xp0OneMv2c/g/rfxE8X6G+v/ABV+"
    "IvmXXijVhBcPb6JarbP9n0/fA/mxLxHAsm9VRnZ3l/dJu9x+BXwn+Iug+MtT8TeMfEDW4mAtT4etr5r20kiW1sokdfkiigCy293I"
    "qRRL8t38392vc9e1/T/C+h3+savdw6dpVhC9zc3Vw21IolXc7t9KAPz+8O+Kl+CXw1HjPR7i+1jUr/VW0HwR4LS7jv01bUUiSztZ"
    "Yr1Fia5sYIlZER4lVW819zs0UtfVP7LPwHPwL+H80eq3v9t+OfEFw2r+J9ccfPd30nzP83/PNPur/wB9Y+avLP2ZvDWoftCfEqX9"
    "oXxLYyafoEcMum/DrQrhdn2PTc7Xv2X+CW4/9A/vLsNfYFABRRXzx+0l+0Ze+A7qx+Hvw6s4vEvxi8QKV03Ss7otNjPDX15j/VwJ"
    "975vvYoA4j4z3S/tMftN+E/hFp5+0eEfAs8XivxpNGf3TXK/8eFgf9pmzK6f3f8AaWm+OPiZ4j8OeMZ/jMuo+Fbf4cS32neErS31"
    "1JEuLmze9WK7u4J/NWKJXd3f50dXitIn+Wuf1bwzb/s+fA/xF4A0DxLdHxbqX+leM/HSokt3NqN58qRRbnXddXDuixJvXyom81nX"
    "5d/pGg+NNZh+Ingz4V6z8MLDxWND0G1l1jxFoqWqado11KrRJFFbyvvij8pJunzbCu1NrUAdF+zH4d8CaTpviKT4baDb6b4ON6sF"
    "lq0NzLcLqqxJ87wvIzZtkd3ii2ts+SXaNuN3A/sjv/wt343fGv42f67StQ1GPwp4cmB+VrCx+WWVP9iWX5v+A10X7ZnxQ1Xwr4F0"
    "/wCHngYq/wASvH8v9iaHbxcG1ibi4vG/upFF/F/CxWvXPg/8MtJ+DXwz8M+CNCTGmaJYraxyMu1pWHLyt/tu5Z2/2noA7aiiigAo"
    "oooAKKKKACiiigAooooAKKKKACvjPxJDP+wv8XNS8YW0Er/AbxpfpLr8NujMvhjVJcL9tVFH/HtKdqv/AHWx/sI32ZWVrmh6f4l0"
    "W90fVbSHUNNvoWt7mzuIw8U8bLtZGU/wkGgD5b+MX7PfhuRvHPxH1/Xtb1/4V/YT4gf4d+F3eKy1GWOHzZbiQrLidpfv7V8pG4Z9"
    "53M2L4c/aC0T4b+OtA0Pwl4O1XS9V15I7i++GWlW0V0kFm0W5dXspbV3gWPavzJuXzeqoku/zZ4Y/FX7AdxMiQah4z/Z2kk3Ky77"
    "rUvB+772V+9PZ/8Ajyf+h53xkt/h5pvh3wpafBbwXoaeJviNqytonjvSIPsun6Vc7flupL2L/lrt3qturfvSXRl+dlYA7jxv8C4f"
    "iBrifGX4DeKrTwn4/nUJdT7HOm68qn5rfUrf7yyL8ybtvmx+m5V2aXgf9sLSodeg8HfGHSZfhD49fCR22szL/ZmoYzl7K9/1ci9P"
    "lYq2W2/NXcfBX9nvRPgvHNc22o6vrPiTUS8uua1f3kpbVrpyrNcSxbvL3fKFXavyp8td14x8C6B8QtBn0TxLomn6/pE3MtlqVqs8"
    "TY/2W/i/2qANxJFkRWVtyt91lrg/jf8AB3Qvjv8ADTWfBfiBWWy1GP8Ad3EP+ttp1+aKeP8A20b5vwryKL9jXUPh3k/Bj4p+J/hp"
    "Av8AqtCu2/tvRUHfba3PzR/8Blqwus/tV+DVIuvDXw5+JNsvCSaZqd1o13J7ukqSxKf91qAJf2TPi/r/AIht9c+FnxHkJ+K3gR1t"
    "dSkZsf2vZ/8ALvqEX95ZV27v9r723cFrjfjhe3H7WnxpX4GaFdTL8P8Aw5LFqPxC1O3JEdy33rfSEb+85XdJt+7t67lZa81/aCsv"
    "2i/Hnijwv4+8F/Ai98D/ABO0AS2kWs2/ijTL+2urOVW8y3mTejOu75k3LhWzXW/s9/8AC6/gf8NrLwzon7OU9zqU0r32sa9r3jew"
    "jl1G/l+ae4l8vzW+Zxx/shaAPtHT9PttJsYLOzhjtbW3jWGGCNdqRqoCqqqOi1W8ReJNJ8J6LcatrmqWmjaXbJvnvdQnWCGJfVnb"
    "5a8Ck0n9qjx5GYrnXvh78KbKUcTaRZ3GuajF/wB//Kg/8dapdF/Yj8GXms2mu/ErV9e+MviCB/Njl8Y3n2iyt3P3vIslCwIn+zsa"
    "gDn9U/aV8ZfH65l0D9nvSRJphbybz4n+ILd4tIs/7xs4m+a8lXn/AKZbtu7cjVynwxbwZ8L/ABla+BPBfiGa78eeNmmn1f4r+IYH"
    "lbWZoivmxWUzL5U8y7/kiVvKi28+a6ujesftCeA4rjwPq9x4hsvFnjbwzFFstvBPgomwXykTdtl8uVJZfuH5d+37ipFu+94pb+C9"
    "P+J3wJ03Ttc07xZ8SNS0PR5RFrWrXM+g6fpFzudoJYpbxbdnniDxp9q8p2VIH+6zPE4A+++D+t/E/wAdSeDb3wLfQ+HdH8Surf8A"
    "CR2KzaRPpP8Ay3vfPdt9zqN6x+WVfnh/vRfP5v0Z4k1r4dfsjfCTUtaltbfwz4a0/dK0VooaW7nb7qru+aWZ+F+Y5P0FfPHhH9oP"
    "XvgneRar44+Klx8Q/AMehR6bYzJokKXXiLXBKq7NJWL97dLt+VpW3R73T5/v7O5+GPwT8XfG3x5pnxb+N9oNOfTW8/wn8PVl8630"
    "Xut1df8APW8/RP8Ae4QAv/szfC/xL4h8Yat8ePilaiz8beIoPsui6FNlh4b0nO5Lcf8ATZ87pW/9A+Za+nqKKACiiigAooooAKKK"
    "KACiiigAooooAKKKKACiiigCJ41kRlZdyt95Wr5R+IX7Hd/4VuNT1z4HahZ+FpNQuBdap4H1YM/hzVZUdGV1jX5rOfciMstvt5Rf"
    "u9a+s6KAPmvwn+2hoOma1B4Y+Lmi3vwZ8YSfKkPiBlbTLxv71rfr+6kX/e2tX0bbzxXkMc0MqywyLuV1bcrLWX4r8I6J460O40bx"
    "DpFlrulXC7ZbLUbdJ4X/AN5GHNeBN+xPYeCZ3u/g5498UfCGZjvGm6fc/wBoaOz/AN97C53L/wB8MtAH0xRXzMmr/tT+Ady3Wh+A"
    "PixYxfKj6fez6DqE3u6yrLAG/wB1ql/4at8a6Cu3xX+zp8SrC5P3v+EfSy1qJf8AgcVx/wCy0AfSlFfN0f7cHh6OPN78Nvizp0n9"
    "268DXob/AMdU0sn7bOlzLnTvhN8YNXyv/Lp4IulH/fUuxaAPpCivmlf2kvix4jUHwl+zd4rkRm2Cbxdq9hoyr/tMm+V/++VqNvDH"
    "7UXxE+TVvGHgv4Taa/IXwzp8us6gF/uPLc7IlP8AtKjUAe++LvGmg+AdDn1jxHrNjoGkwf6291K5WCJf+BNxXyX4w+KKftUa9Zj4"
    "O/C6w8bCzHlw/EXxvYPFoWnHd9+1ilXfcyr/ALCr/D8xWvRvC/7Evw7sdbt/EHjNtZ+K/iaE5j1bx5fvqPl+0Vv/AKhF/wCAfLxX"
    "0JDBHawpFGixxIu1VUYVVoA8N+Df7KulfD/xFJ428V61efEb4o3EflzeKtaRc2yf88rK3UbLWLr8q/N8zfNzXvNFFABRRRQAUUUU"
    "AFFFFABRRRQAUUUUAFFFFABRRXzH8bPj9441L4vH4MfB7TtPPjOPTl1XWPEmvBvsGi2rPsXbEvzTztlSq/d+ZPvfPsAPpyolmjdi"
    "qurOv8O6vl61/YWsvGSm6+LvxL8b/FC8l5ntJdVl0vSv+2dpasmxf+B1wv7Qn7G37OvwP+Cvi3x3B4Em0m70Cxa6tLzR9avILtJ/"
    "uxFZfN+9vZfvbutAH3BRXwz+zr+098UPAvwT8J6b8RPgv8UvFuvQ2ZafXNMsbe9+2Rs7NE/M6tuETIvzfM22vSYf27NNh/5CXwX+"
    "NOix95L7wTKVX8YnegD6dor5i/4eGfCq3H/EwtPGejH/AKfvCV+v/oMTU1v+Ck37PMfF146urF/7t14e1OI/rb0AfT9FfNUP/BRz"
    "9nCYfL8UdNB/2rS6X/2lV6P/AIKCfs7yL8vxW0If73mr/wCyUAfQ9FfP/wDw3t+z3/0Vjw7/AN/2/wDiaVv2+P2e1Xd/wtjw/wD9"
    "/X/+JoA9/or51n/4KDfs7xdfitoh/wB1ZW/9krOn/wCCkX7OESf8lPs3/wBmHT712/8AHYaAPpuivmJf+CkPwCmx9k8WalqJb7v2"
    "Twzqj7v/ACXob/goF8ObjA0zw/8AEDWS33fsPg+/bd/30i0AfTtFfMX/AA3JFdSbdP8AgN8btS/uuvg0xJ/31LKtfOn7YHx68WeP"
    "Jfh2fEfw6+Ifw/8Ag5FrKweLmvbiPTXv4p3VIkZoJXbYuHdk/iz/AA/K1AH6RrNG7FVdWdf4d1S18wv/AME3f2dFjVYPhytpNHyl"
    "xa6vfxSq397etxuqhffsr/Ef4WxvqHwV+MWvQzRfOvhbx9O2r6TPjpErt+/t16fMrMaAPq2ivEf2Y/2hJfjnoviG31nQ38L+NvCu"
    "pNoviHRzL5sUN0n8cMn8cT8lf68M3t1ABRRRQAUUUUAFFFFABRRRQAV8B/tyfCHRT+1V8EfHOrXmr6NoviKdvBuq6loOotYXME77"
    "3snEq9PnZ93+ylffleDftr/CC8+N37Nni7QdGRpPElrGmqaN5PMv2y2YSoqf7T7Wj/4HQB82+MLKy/ZG/bJ+D39qfFrxhL4N1bS9"
    "Yk1WXxx4okurT91bt5K5l2r99k+X+9sqr+2F+1N8Ov2ovCfhP4PfD3W7rX5fF/izS9Ou7qHTriG1NqtwGl2SyoquVdYj8m6sOy+L"
    "3hX9q79sT9m621fRkv8AXNK0TVF8UeG9b0h1GnX4tN+x4p02/LKm5G5/g+61faXxu/Zj8FfHqPw3/b7arptz4blln0q70G/exmtW"
    "cIHKsn+4tAHrMcSwxrHEqoirtVV/hr5m8H/Ff4k3v7R2pya14b12x+GOpR3+m6L9ptIkhEtnsdbhvm89PP2Xv+tRF2pb7d2+hP2H"
    "Xs23ad8fvjbYn+GNvF3nov8AwGWJqlX9lb4m6b/yCP2mvHUR7DUrGwvf/Q4qAM/SP2mJv+GbvAHiW78T2L+Jhd+HLfxVcXYiiWzN"
    "zcW/23zl2qsX7r7R/d21F8Z/2kLqw+K3gnS/C3jfRrLwvrGkSah/aH2zTvKun+1JEnly3Mqq/wDH8kWWrRb4A/tE27f6H+1NM6D7"
    "q33gHTZf/HlZKRvg/wDtOwLtHx98N6jt+79r8CxJu/75uKALHh/x98RfGP7Q/iTRrTSLC8+Hui6+2l6hdXFlCqQQ/wBlW8+5Z/tH"
    "mvP586Ls8jZsf7/y15boHxu+JV1+zvrvxH1nwZ4fS0/sPS9Q0e8u9Gjt0lurm42yxeV9tbzYlieFlZzB8zf98+kw+Af2rLLcYfid"
    "8Ornc25vO8Lzpu/75lqvP4H/AGqv7NOnv4k+EF5Ybdn2SfQL0RMv93Z5u2gClrfjjXfDPwbsfG9t4T8KeJp11Mabfad9gs7Vm891"
    "t7VkeC9uol23Etvu/et+6dvlVl+b3zwv4T0j+y7aHUNN8PXeu2kUUOpPptkqQpceWrsFU7mRTuDKjNna614Vb+Af2pLOw+w2uo/B"
    "O2svN837NDoWoLFv3bt+3zfvbhUtl4T/AGsdPutQubXXfg3DPfzrcXTrpep/vJViSLd/rf7kSL/wGgDjvAv7R2v+IPCvjzUL/wAM"
    "aB4Zv9C8JanrGkQXGnH/AInU9rLcRNdJ8/ywRPBEjRff/eo/yqyVeh/aB8c3nwd0nWNIvbPUvEd94l07SHi0q10yWRbedPuIsWoS"
    "xK+77vmyxf7tdGPh/wDtSyrEreKfhHbLGrJH5Ph68farfe+9L/F3p1v8KP2noV2wfE/4eaUm7dtsfCD/APs0tAHKa1+018R7j4L/"
    "AA9m8OWJ174iapeajdahp+naU8sy2VhLKstvcQfN9nuGdrO3l/hSWWXZ8qrXrXhf9oCw1H4uSWN3rDJ4c8QeHtC1XwzDNaFGllup"
    "b1JV3Ku7d+6tdyP9zd71zb/Bn9pu8ZfN/aN0axXv9l8A27/+h3FJH+zv8frjd9v/AGptTdT2svBOmW//AMVQBJ8FPiZ8S9Y+NHid"
    "fE/hzXoPh9rlxdf8I5e3FlEsNn9kleLYfKbzUS4iVJVNwifMjqu7etbv7dngj/hYf7IvxR0lIzNIukPqESp95ntXW5UL+MVYDfsk"
    "eOtUX/ib/tJ/EyT/ALBj2dl/6DBVWb9g3S9Uhlg1v40/GbxBaypsks9R8ZMYZF/iyiRLQBk/BD/got8G/FXgnwjb+JfGS+GfEtzp"
    "lv8Aak162ns4nn8pfNZbh08pl37vm314H+y/8LbH4jfsm6p8XvHnxU+KJitG1S8vrXS/F9xBazQWzy4wn+4v96vvLwX8B/BPgn4U"
    "6b8OLbQrbU/CGnxvFHp2roL1CrOztu83du+Z2/OvzC+DXjUfFD9jnQv2bvBP9oXHjPX/ABPJb6+lnYy+Xo+lG8eWW4lk2eUo2oi7"
    "d3dqAPsn/gmR8K5Phz+y3o+rXcLRav4xuZfEFz5zF32S/Lbjc3zN+5SN+e7tX1zWZoui2Xh/RtP0mwgS3sLCFLa2hT7scaKFRf8A"
    "vnFadABRRRQAUUUUAFFFFABRRRQAUUUUAQeTE0izFFLqu1Xx81T0UUAFFFFABRRRQAUUUUAFFFFABRRRQAUUUUAFFFFABUFvDFDH"
    "thRY0znai7RRRQBPRRRQAUUUUAFFFFAH/9k="
)


class MarkSheetError(Exception):
    """Raised for problems with the uploaded student list."""
    pass


def parse_marks_docx(file) -> pd.DataFrame:
    """
    Reads an uploaded .docx file containing a table with columns:
    Sl.No, Name, APAAR ID, Marks (in that order, but matched by header
    keyword so column order in the source file doesn't matter).
    Column names are matched case-insensitively and don't need to match
    exactly (e.g. 'Sl.no', 'SL NO', 'Marks Obtained' are all accepted).
    Returns a DataFrame with columns: Sl.No, Name, APAAR ID, Marks (float).
    """
    try:
        doc = Document(file)
    except Exception as exc:
        raise MarkSheetError(
            "Could not open the uploaded file. Please make sure it is a valid .docx file."
        ) from exc

    if not doc.tables:
        raise MarkSheetError(
            "No table was found in the uploaded document. "
            "Please upload a .docx file with a table containing Sl.No, Name, APAAR ID and Marks columns."
        )

    table = doc.tables[0]
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]

    if len(rows) < 2:
        raise MarkSheetError("The table in the uploaded document has no data rows.")

    header = [h.lower() for h in rows[0]]

    def find_col(keywords):
        for i, h in enumerate(header):
            if any(k in h for k in keywords):
                return i
        return None

    idx_sl = find_col(["sl"])
    idx_name = find_col(["name"])
    idx_apaar = find_col(["apaar"])
    idx_marks = find_col(["mark"])

    missing = []
    if idx_sl is None:
        missing.append("Sl.No")
    if idx_name is None:
        missing.append("Name")
    if idx_apaar is None:
        missing.append("APAAR ID")
    if idx_marks is None:
        missing.append("Marks")
    if missing:
        raise MarkSheetError(
            "Could not find the following required column(s) in the uploaded table: "
            + ", ".join(missing)
            + ". Please check the column headers in the uploaded file."
        )

    records = []
    for r in rows[1:]:
        if not any(cell.strip() for cell in r):
            continue  # skip fully blank rows
        name = r[idx_name].strip()
        if not name:
            continue
        raw_marks = r[idx_marks].strip()
        try:
            marks_val = float(raw_marks)
        except ValueError:
            raise MarkSheetError(
                f"Could not read marks for '{name}' (value: '{raw_marks}'). "
                "Marks must be numeric."
            )
        records.append(
            {
                "Sl.No": r[idx_sl].strip(),
                "Name": name,
                "APAAR ID": r[idx_apaar].strip(),
                "Marks": marks_val,
            }
        )

    if not records:
        raise MarkSheetError("No valid student rows were found in the uploaded table.")

    return pd.DataFrame(records)


def compute_marks(df: pd.DataFrame, max_test_marks: float, sa_component: float) -> pd.DataFrame:
    """
    Adds two columns to the DataFrame:
    - Marks_25: raw Marks scaled to out of 25, rounded to 1 decimal
    - SA: (Marks_25 / 25) * sa_component, rounded to 1 decimal
    """
    if max_test_marks <= 0:
        raise MarkSheetError("Maximum Test Marks must be greater than 0.")

    out = df.copy()
    out["Marks_25"] = (out["Marks"] / max_test_marks * 25).round(1)
    out["SA"] = (out["Marks_25"] / 25 * sa_component).round(1)
    return out


def _set_no_borders(table):
    """Remove all borders from a python-docx table (used for the details table)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def _get_logo_stream():
    """Decodes the embedded base64 logo into an in-memory stream for add_picture()."""
    try:
        return io.BytesIO(base64.b64decode(LOGO_BASE64))
    except Exception:
        return None


def generate_marksheet(
    department: str,
    course_code: str,
    course_name: str,
    max_test_marks: float,
    sa_component: float,
    df: pd.DataFrame,
    college_name: str = "Sanatana Dharma College, Alappuzha",
    exam_title: str = "Internal Examination, September 2026",
) -> io.BytesIO:
    """
    Builds the formatted mark sheet .docx (matching the college template layout)
    and returns it as an in-memory BytesIO buffer, ready for download.
    """
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    # --- Header: logo + college name + exam title ---
    logo_stream = _get_logo_stream()
    if logo_stream is not None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(logo_stream, width=Inches(0.9))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(college_name)
    run.bold = True
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(exam_title)
    run.font.size = Pt(12)

    doc.add_paragraph()  # spacer

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Mark Sheet")
    run.bold = True
    run.underline = True
    run.font.size = Pt(13)

    doc.add_paragraph()  # spacer

    # --- Details: Department / Course Code / Course Name ---
    details_table = doc.add_table(rows=3, cols=2)
    details_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_no_borders(details_table)
    details_table.columns[0].width = Inches(1.8)
    details_table.columns[1].width = Inches(4.5)

    fields = [
        ("Department", department),
        ("Course Code", course_code),
        ("Course Name", course_name),
    ]
    for i, (label, value) in enumerate(fields):
        left_cell = details_table.cell(i, 0)
        left_cell.text = ""
        left_p = left_cell.paragraphs[0]
        left_run = left_p.add_run(f"{label}\t:")
        left_run.bold = True

        right_cell = details_table.cell(i, 1)
        right_cell.text = str(value) if value else ""

    doc.add_paragraph()  # spacer

    # --- Marks table: Sl.No | Name | APAAR ID | Marks | Marks (out of 25) | SA ---
    marks_table = doc.add_table(rows=1, cols=6)
    marks_table.style = "Table Grid"
    marks_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Inches(0.6), Inches(1.7), Inches(1.4), Inches(0.9), Inches(1.2), Inches(1.2)]
    headers = [
        "Sl.No",
        "Name",
        "APAAR ID",
        "Marks",
        "Marks\n(out of 25)",
        f"SA\n(out of {sa_component:g})",
    ]
    hdr_cells = marks_table.rows[0].cells
    for i, htext in enumerate(headers):
        hdr_cells[i].width = widths[i]
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(htext)
        run.bold = True

    for _, row in df.iterrows():
        cells = marks_table.add_row().cells
        values = [
            str(row["Sl.No"]),
            str(row["Name"]),
            str(row["APAAR ID"]),
            f"{row['Marks']:.1f}",
            f"{row['Marks_25']:.1f}",
            f"{row['SA']:.1f}",
        ]
        for i, val in enumerate(values):
            cells[i].width = widths[i]
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i not in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(val)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
